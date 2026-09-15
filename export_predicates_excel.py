#!/usr/bin/env python3
"""
Export nuPlan predicate Parquet assertions to a compact Excel workbook.

Main difference from the previous exporter:
- By default, only "relevant" columns are written.
- Bookkeeping/provenance columns that make Excel unnecessarily wide are removed.
- Columns that are completely empty for the selected predicate family/families
  are not written.
- Predicate-specific measurements are retained automatically.
- Use --columns all if you ever want the old full-column behavior.

Typical interaction export:

    python export_predicates_excel.py \
      --output-dir outputs_python/paper_release \
      --families interaction

Output:
    <output-dir>/predicates_by_family.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set

try:
    import pandas as pd
except ImportError as exc:
    raise SystemExit("ERROR: pandas is required.") from exc

try:
    import pyarrow.parquet as pq
except ImportError as exc:
    raise SystemExit(
        "ERROR: pyarrow is required. Install it with: pip install pyarrow"
    ) from exc

try:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
except ImportError as exc:
    raise SystemExit(
        "ERROR: openpyxl is required. Install it with: pip install openpyxl"
    ) from exc


EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_DATA_ROWS = EXCEL_MAX_ROWS - 1
EXCEL_MAX_TEXT_LENGTH = 32_767


# ---------------------------------------------------------------------------
# Compact Excel policy
# ---------------------------------------------------------------------------

# These columns are useful for understanding WHERE/WHEN/WHO a predicate
# assertion belongs to. They are kept first when they exist.
CORE_COLUMNS = [
    "predicate_id",
    "subject_id",
    "object_id",
    "valid_time_us",
    "end_time_us",
    "scenario_token",
    "prov_scenario_token",
    "scenario_type",
    "prov_scenario_type",
]

# These are normally implementation/provenance bookkeeping rather than
# information needed to inspect predicate detections in Excel.
#
# Predicate-specific metric columns are NOT listed here, so fields such as:
#   map_bumper_gap_m
#   bumper_headway_s
#   condition_duration_s
#   subject_forward_speed_mps
#   travel_direction_difference_rad
#   ...
# are kept automatically if they contain data.
NOISE_COLUMNS = {
    "assertion_kind",
    "log_name",
    "timestamp_us",
    "rule_id",
    "rule_version",
    "extractor_version",
    "generated_at",
    "value",
    "value_type",
    "evidence",
}

# Also suppress common low-level export/provenance fields by prefix.
NOISE_PREFIXES = (
    "_",
    "generated_",
    "extractor_",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Export predicate Parquet batches to a compact Excel workbook, "
            "grouped by predicate family."
        )
    )

    parser.add_argument(
        "--output-dir",
        required=True,
        type=Path,
        help="Predicate output directory.",
    )

    parser.add_argument(
        "--excel",
        type=Path,
        default=None,
        help=(
            "Destination .xlsx file. "
            "Default: <output-dir>/predicates_by_family.xlsx"
        ),
    )

    parser.add_argument(
        "--families",
        default="all",
        help=(
            "Comma-separated families to export, or 'all'. "
            "Example: interaction"
        ),
    )

    parser.add_argument(
        "--batch-rows",
        type=int,
        default=25_000,
        help="Rows read from Parquet per streaming batch. Default: 25000",
    )

    parser.add_argument(
        "--rows-per-sheet",
        type=int,
        default=EXCEL_MAX_DATA_ROWS,
        help=(
            "Maximum data rows per worksheet, excluding the header. "
            f"Default: {EXCEL_MAX_DATA_ROWS}"
        ),
    )

    parser.add_argument(
        "--no-position-values",
        action="store_true",
        help=(
            "Do not split spatial literal/value predicates into "
            "'position_values'."
        ),
    )

    parser.add_argument(
        "--columns",
        choices=["relevant", "all"],
        default="relevant",
        help=(
            "Column policy. 'relevant' removes bookkeeping and completely "
            "empty columns. 'all' keeps the original full schema. "
            "Default: relevant"
        ),
    )

    parser.add_argument(
        "--include-evidence",
        action="store_true",
        help=(
            "Keep the raw evidence column in relevant-column mode. "
            "Normally omitted because it can make Excel very large."
        ),
    )

    return parser.parse_args()


def find_parquet_files(output_dir: Path) -> List[Path]:
    patterns = [
        "workers/worker_*/batches/batch_*.parquet",
        "batches/batch_*.parquet",
        "**/batches/batch_*.parquet",
    ]

    found: Dict[str, Path] = {}

    for pattern in patterns:
        for path in output_dir.glob(pattern):
            if path.is_file():
                found[str(path.resolve())] = path.resolve()

    return sorted(found.values(), key=lambda p: str(p))


def find_definition_files(output_dir: Path) -> List[Path]:
    preferred = output_dir / "predicate_definitions.csv"
    files: List[Path] = []

    if preferred.is_file():
        files.append(preferred.resolve())

    for path in output_dir.glob("workers/worker_*/predicate_definitions.csv"):
        if path.is_file():
            files.append(path.resolve())

    if not files:
        for path in output_dir.rglob("predicate_definitions.csv"):
            if path.is_file():
                files.append(path.resolve())

    seen: Set[str] = set()
    result: List[Path] = []

    for path in files:
        key = str(path)
        if key not in seen:
            result.append(path)
            seen.add(key)

    return result


def load_definition_map(
    definition_files: List[Path],
) -> Dict[str, Dict[str, str]]:
    if not definition_files:
        raise FileNotFoundError(
            "No predicate_definitions.csv was found under the output directory."
        )

    mapping: Dict[str, Dict[str, str]] = {}

    for path in definition_files:
        df = pd.read_csv(path)

        required = {"predicate_id", "category"}
        missing = required - set(df.columns)

        if missing:
            raise ValueError(
                f"{path} is missing required columns: {sorted(missing)}"
            )

        for _, row in df.iterrows():
            predicate_id = str(row["predicate_id"])
            category = str(row["category"])

            value_type = (
                str(row["value_type"])
                if "value_type" in df.columns and pd.notna(row["value_type"])
                else ""
            )

            incoming = {
                "category": category,
                "value_type": value_type,
            }

            current = mapping.get(predicate_id)

            if current is None:
                mapping[predicate_id] = incoming
            elif current != incoming:
                print(
                    "WARNING: conflicting predicate definition for "
                    f"{predicate_id}; keeping first definition {current}, "
                    f"ignoring {incoming} from {path}",
                    file=sys.stderr,
                )

    return mapping


def family_for_predicate(
    predicate_id: str,
    definition_map: Dict[str, Dict[str, str]],
    split_position_values: bool,
) -> str:
    definition = definition_map.get(predicate_id)

    if definition is None:
        return "unknown"

    category = (definition.get("category") or "unknown").strip() or "unknown"
    value_type = (definition.get("value_type") or "").strip().lower()

    if (
        split_position_values
        and category == "spatial"
        and value_type != "entity"
    ):
        return "position_values"

    return category


def discover_columns(parquet_files: List[Path]) -> List[str]:
    union: List[str] = []
    seen: Set[str] = set()

    for path in parquet_files:
        schema = pq.ParquetFile(path).schema_arrow

        for name in schema.names:
            if name not in seen:
                union.append(name)
                seen.add(name)

    return union


def is_noise_column(name: str, include_evidence: bool) -> bool:
    if name == "evidence" and include_evidence:
        return False

    if name in NOISE_COLUMNS:
        return True

    return any(name.startswith(prefix) for prefix in NOISE_PREFIXES)


def has_real_value(series: pd.Series) -> bool:
    """
    True if at least one value is useful/non-empty.

    Empty strings, NaN, None, and empty containers do not count.
    """
    for value in series:
        if value is None:
            continue

        try:
            missing = pd.isna(value)
            if isinstance(missing, bool) and missing:
                continue
        except Exception:
            pass

        if isinstance(value, str):
            if value.strip():
                return True
            continue

        if isinstance(value, (list, tuple, set, dict)):
            if len(value) > 0:
                return True
            continue

        return True

    return False


def determine_relevant_columns(
    parquet_files: List[Path],
    definition_map: Dict[str, Dict[str, str]],
    requested_families: Optional[Set[str]],
    split_position_values: bool,
    batch_rows: int,
    all_columns: List[str],
    include_evidence: bool,
) -> List[str]:
    """
    First pass over the selected rows.

    Keep:
    1. useful identity/context columns;
    2. every non-bookkeeping column that actually contains at least one value
       for the requested family/families.

    This automatically keeps predicate-specific measurements without needing
    a hard-coded list for follows/overtakes/merge/etc.
    """
    candidates = [
        c
        for c in all_columns
        if not is_noise_column(c, include_evidence=include_evidence)
    ]

    nonempty: Set[str] = set()

    # predicate_id is required for family filtering.
    read_columns = list(dict.fromkeys(["predicate_id"] + candidates))

    print("Scanning selected rows to find non-empty relevant columns...")

    selected_rows = 0

    for file_index, parquet_path in enumerate(parquet_files, start=1):
        parquet_file = pq.ParquetFile(parquet_path)

        available = set(parquet_file.schema_arrow.names)
        file_read_columns = [c for c in read_columns if c in available]

        for batch in parquet_file.iter_batches(
            batch_size=batch_rows,
            columns=file_read_columns,
        ):
            frame = batch.to_pandas()

            if "predicate_id" not in frame.columns:
                continue

            families = frame["predicate_id"].astype(str).map(
                lambda pid: family_for_predicate(
                    pid,
                    definition_map,
                    split_position_values=split_position_values,
                )
            )

            if requested_families is None:
                mask = pd.Series(True, index=frame.index)
            else:
                mask = families.isin(requested_families)

            selected = frame.loc[mask]

            if selected.empty:
                continue

            selected_rows += len(selected)

            for column in selected.columns:
                if column == "predicate_id":
                    nonempty.add(column)
                    continue

                if column in nonempty:
                    continue

                if has_real_value(selected[column]):
                    nonempty.add(column)

        print(
            f"  [{file_index}/{len(parquet_files)}] scanned "
            f"{parquet_path.name}",
            end="\r",
            flush=True,
        )

    print()
    print("Selected rows seen during column scan:", f"{selected_rows:,}")

    # Stable readable order:
    # core columns first, then predicate-specific measurements in source order.
    ordered: List[str] = []

    for column in CORE_COLUMNS:
        if column in nonempty and column not in ordered:
            ordered.append(column)

    for column in all_columns:
        if column in nonempty and column not in ordered:
            ordered.append(column)

    return ordered


_INVALID_SHEET_CHARS = re.compile(r"[\[\]\:\*\?\/\\]")


def sanitize_sheet_name(name: str) -> str:
    name = _INVALID_SHEET_CHARS.sub("_", str(name)).strip()
    name = name or "unknown"
    return name[:31]


def split_sheet_name(family: str, part: int) -> str:
    suffix = f"_{part}"
    base = sanitize_sheet_name(family)

    if len(base) + len(suffix) > 31:
        base = base[: 31 - len(suffix)]

    return f"{base}{suffix}"


def excel_value(value: Any) -> Any:
    if value is None:
        return None

    try:
        missing = pd.isna(value)
        if isinstance(missing, bool) and missing:
            return None
    except Exception:
        pass

    if hasattr(value, "item") and not isinstance(
        value,
        (str, bytes, bytearray),
    ):
        try:
            value = value.item()
        except Exception:
            pass

    if isinstance(value, (datetime, date, int, float, bool)):
        if isinstance(value, float) and (
            math.isnan(value) or math.isinf(value)
        ):
            return None if math.isnan(value) else str(value)
        return value

    if isinstance(value, (bytes, bytearray)):
        value = value.hex()

    elif isinstance(value, (dict, list, tuple, set)):
        try:
            value = json.dumps(
                value,
                ensure_ascii=False,
                sort_keys=True,
                default=str,
            )
        except Exception:
            value = str(value)

    else:
        value = str(value)

    if len(value) > EXCEL_MAX_TEXT_LENGTH:
        suffix = "... [TRUNCATED FOR EXCEL]"
        value = value[: EXCEL_MAX_TEXT_LENGTH - len(suffix)] + suffix

    return value


class FamilySheetWriter:
    def __init__(
        self,
        workbook: Workbook,
        columns: List[str],
        rows_per_sheet: int,
    ) -> None:
        self.workbook = workbook
        self.columns = columns
        self.rows_per_sheet = rows_per_sheet
        self.states: Dict[str, Dict[str, Any]] = {}

    def _create_sheet(self, family: str, part: int):
        if part == 1:
            title = sanitize_sheet_name(family)
        else:
            state = self.states.get(family)

            if state is not None and part == 2:
                first_ws = state["worksheets"][0]
                first_ws.title = split_sheet_name(family, 1)

            title = split_sheet_name(family, part)

        ws = self.workbook.create_sheet(title=title)

        header = []

        for column in self.columns:
            cell = WriteOnlyCell(ws, value=column)
            cell.font = Font(bold=True)
            header.append(cell)

        ws.append(header)
        return ws

    def append(self, family: str, record: Dict[str, Any]) -> None:
        state = self.states.get(family)

        if state is None:
            ws = self._create_sheet(family, 1)

            state = {
                "part": 1,
                "rows_in_part": 0,
                "total_rows": 0,
                "worksheet": ws,
                "worksheets": [ws],
            }

            self.states[family] = state

        if state["rows_in_part"] >= self.rows_per_sheet:
            state["part"] += 1

            ws = self._create_sheet(family, state["part"])

            state["worksheet"] = ws
            state["worksheets"].append(ws)
            state["rows_in_part"] = 0

        row = [
            excel_value(record.get(column))
            for column in self.columns
        ]

        state["worksheet"].append(row)
        state["rows_in_part"] += 1
        state["total_rows"] += 1


def main() -> int:
    args = parse_args()

    output_dir = args.output_dir.expanduser().resolve()

    if not output_dir.is_dir():
        raise SystemExit(
            f"ERROR: output directory does not exist: {output_dir}"
        )

    if args.batch_rows <= 0:
        raise SystemExit("ERROR: --batch-rows must be > 0")

    if not (1 <= args.rows_per_sheet <= EXCEL_MAX_DATA_ROWS):
        raise SystemExit(
            "ERROR: --rows-per-sheet must be between 1 and "
            f"{EXCEL_MAX_DATA_ROWS}"
        )

    excel_path = (
        args.excel.expanduser().resolve()
        if args.excel is not None
        else output_dir / "predicates_by_family.xlsx"
    )

    excel_path.parent.mkdir(parents=True, exist_ok=True)

    parquet_files = find_parquet_files(output_dir)

    if not parquet_files:
        raise SystemExit(
            "ERROR: no batch_*.parquet files were found under "
            f"{output_dir}"
        )

    definition_files = find_definition_files(output_dir)
    definition_map = load_definition_map(definition_files)
    all_columns = discover_columns(parquet_files)

    if "predicate_id" not in all_columns:
        raise SystemExit(
            "ERROR: Parquet assertion files do not contain "
            "'predicate_id'."
        )

    requested_families: Optional[Set[str]]

    if args.families.strip().lower() == "all":
        requested_families = None
    else:
        requested_families = {
            item.strip()
            for item in args.families.split(",")
            if item.strip()
        }

        if not requested_families:
            raise SystemExit(
                "ERROR: --families was provided but no family names "
                "were parsed."
            )

    split_position_values = not args.no_position_values

    if args.columns == "all":
        columns = list(all_columns)
    else:
        columns = determine_relevant_columns(
            parquet_files=parquet_files,
            definition_map=definition_map,
            requested_families=requested_families,
            split_position_values=split_position_values,
            batch_rows=args.batch_rows,
            all_columns=all_columns,
            include_evidence=args.include_evidence,
        )

    if not columns:
        raise SystemExit(
            "ERROR: no relevant columns were found for the selected families."
        )

    print()
    print("Predicate output directory:", output_dir)
    print("Parquet files found:", len(parquet_files))
    print("Predicate definition files found:", len(definition_files))
    print("Excel output:", excel_path)
    print("Column mode:", args.columns)
    print("Columns written:", len(columns))
    print("Selected columns:")
    for column in columns:
        print("  -", column)

    print(
        "Families:",
        "all"
        if requested_families is None
        else ",".join(sorted(requested_families)),
    )

    print()

    workbook = Workbook(write_only=True)

    writer = FamilySheetWriter(
        workbook=workbook,
        columns=columns,
        rows_per_sheet=args.rows_per_sheet,
    )

    predicate_counts: Counter[str] = Counter()
    family_counts: Counter[str] = Counter()
    unknown_predicates: Counter[str] = Counter()

    processed_rows = 0
    written_rows = 0

    # Only read predicate_id + columns that will actually be written.
    wanted_columns = list(dict.fromkeys(["predicate_id"] + columns))

    for file_index, parquet_path in enumerate(parquet_files, start=1):
        parquet_file = pq.ParquetFile(parquet_path)
        file_rows = parquet_file.metadata.num_rows

        available = set(parquet_file.schema_arrow.names)
        file_columns = [c for c in wanted_columns if c in available]

        print(
            f"[{file_index}/{len(parquet_files)}] "
            f"{parquet_path} ({file_rows:,} rows)"
        )

        for batch in parquet_file.iter_batches(
            batch_size=args.batch_rows,
            columns=file_columns,
        ):
            frame = batch.to_pandas()

            for row in frame.to_dict(orient="records"):
                processed_rows += 1

                predicate_id = str(row.get("predicate_id", ""))

                family = family_for_predicate(
                    predicate_id,
                    definition_map,
                    split_position_values=split_position_values,
                )

                predicate_counts[predicate_id] += 1
                family_counts[family] += 1

                if family == "unknown":
                    unknown_predicates[predicate_id] += 1

                if (
                    requested_families is not None
                    and family not in requested_families
                ):
                    continue

                writer.append(family, row)
                written_rows += 1

            print(
                f"    processed {processed_rows:,} rows | "
                f"written {written_rows:,}",
                end="\r",
                flush=True,
            )

        print()

    if written_rows == 0:
        raise SystemExit(
            "ERROR: no rows matched the requested families; "
            "Excel file was not written."
        )

    # ------------------------------------------------------------------
    # Small summary sheets
    # ------------------------------------------------------------------

    summary_ws = workbook.create_sheet(title="_summary")

    header = []

    for value in ["family", "assertion_count", "excel_sheet_parts"]:
        cell = WriteOnlyCell(summary_ws, value=value)
        cell.font = Font(bold=True)
        header.append(cell)

    summary_ws.append(header)

    for family, count in sorted(
        family_counts.items(),
        key=lambda item: (-item[1], item[0]),
    ):
        if (
            requested_families is not None
            and family not in requested_families
        ):
            continue

        state = writer.states.get(family)
        parts = state["part"] if state is not None else 0

        summary_ws.append([family, count, parts])

    predicates_ws = workbook.create_sheet(
        title="_predicate_counts"
    )

    header = []

    for value in ["predicate_id", "family", "assertion_count"]:
        cell = WriteOnlyCell(predicates_ws, value=value)
        cell.font = Font(bold=True)
        header.append(cell)

    predicates_ws.append(header)

    for predicate_id, count in sorted(
        predicate_counts.items(),
        key=lambda item: (-item[1], item[0]),
    ):
        family = family_for_predicate(
            predicate_id,
            definition_map,
            split_position_values=split_position_values,
        )

        if (
            requested_families is not None
            and family not in requested_families
        ):
            continue

        predicates_ws.append(
            [predicate_id, family, count]
        )

    info_ws = workbook.create_sheet(title="_export_info")

    info_rows = [
        ("source_output_dir", str(output_dir)),
        ("parquet_file_count", len(parquet_files)),
        ("rows_processed", processed_rows),
        ("rows_written", written_rows),
        ("column_mode", args.columns),
        ("columns_written", len(columns)),
        ("include_evidence", args.include_evidence),
        ("rows_per_sheet", args.rows_per_sheet),
        (
            "families",
            "all"
            if requested_families is None
            else ",".join(sorted(requested_families)),
        ),
        ("definition_file", str(definition_files[0])),
    ]

    for key, value in info_rows:
        info_ws.append([key, value])

    if unknown_predicates:
        unknown_ws = workbook.create_sheet(
            title="_unknown_predicates"
        )

        unknown_ws.append(
            ["predicate_id", "assertion_count"]
        )

        for predicate_id, count in sorted(
            unknown_predicates.items(),
            key=lambda item: (-item[1], item[0]),
        ):
            unknown_ws.append([predicate_id, count])

    print()
    print("Saving Excel workbook...")

    workbook.save(excel_path)

    print()
    print("DONE")
    print("Excel file:", excel_path)
    print("Assertions processed:", f"{processed_rows:,}")
    print("Assertions written:", f"{written_rows:,}")
    print("Families written:", len(writer.states))

    for family, state in sorted(writer.states.items()):
        print(
            f"  {family}: {state['total_rows']:,} rows "
            f"({state['part']} sheet(s))"
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())